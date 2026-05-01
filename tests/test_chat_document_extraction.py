"""Unit tests for chat document extraction (txt/md/pdf/docx/spreadsheet) and LLM section budgets."""
from __future__ import annotations

import base64
from io import BytesIO

from docx import Document

from src.ham.chat_document_extraction import (
    MAX_EXTRACTED_CHARS_PER_FILE,
    DocumentExtractionResult,
    build_document_llm_sections,
    extract_document_bytes,
    format_document_block_for_llm,
)

# Small one-page PDF with visible text "Hello PDF Line" (generated once via fpdf2).
_HELLO_PDF_B64 = (
    "JVBERi0xLjMKJenr8b8KMSAwIG9iago8PAovQ291bnQgMQovS2lkcyBbMyAwIFJdCi9NZWRpYUJveCBbMCAwIDU5NS4yOCA4NDEuODldCi9UeXBlIC9QYWdlcwo+PgplbmRvYmoKMiAwIG9iago8PAovT3BlbkFjdGlvbiBbMyAwIFIgL0ZpdEggbnVsbF0KL1BhZ2VMYXlvdXQgL09uZUNvbHVtbgovUGFnZXMgMSAwIFIKL1R5cGUgL0NhdGFsb2cKPj4KZW5kb2JqCjMgMCBvYmoKPDwKL0NvbnRlbnRzIDQgMCBSCi9QYXJlbnQgMSAwIFIKL1Jlc291cmNlcyA2IDAgUgovVHlwZSAvUGFnZQo+PgplbmRvYmoKNCAwIG9iago8PAovRmlsdGVyIC9GbGF0ZURlY29kZQovTGVuZ3RoIDc2Cj4+CnN0cmVhbQp4nDNS8OIy0DM1VyjncgpR0HczVDAy0TMwUAhJU3ANAQkZWegZmyqYW5jqGRkohKQoaHik5uTkKwS4uCn4ZOalaiqEZIFUAgA0OBCpCmVuZHN0cmVhbQplbmRvYmoKNSAwIG9iago8PAovQmFzZUZvbnQgL0hlbHZldGljYQovRW5jb2RpbmcgL1dpbkFuc2lFbmNvZGluZwovU3VidHlwZSAvVHlwZTEKL1R5cGUgL0ZvbnQKPj4KZW5kb2JqCjYgMCBvYmoKPDwKL0ZvbnQgPDwvRjEgNSAwIFI+PgovUHJvY1NldCBbL1BERiAvVGV4dCAvSW1hZ2VCIC9JbWFnZUMgL0ltYWdlSV0KPj4KZW5kb2JqCjcgMCBvYmoKPDwKL0NyZWF0aW9uRGF0ZSAoRDoyMDI2MDUwMTA1MjY1OVopCj4+CmVuZG9iagp4cmVmCjAgOAowMDAwMDAwMDAwIDY1NTM1IGYgCjAwMDAwMDAwMTUgMDAwMDAgbiAKMDAwMDAwMDEwMiAwMDAwMCBuIAowMDAwMDAwMjA1IDAwMDAwIG4gCjAwMDAwMDAyODUgMDAwMDAgbiAKMDAwMDAwMDQzMiAwMDAwMCBuIAowMDAwMDAwNTI5IDAwMDAwIG4gCjAwMDAwMDA2MTYgMDAwMDAgbiAKdHJhaWxlcgo8PAovU2l6ZSA4Ci9Sb290IDIgMCBSCi9JbmZvIDcgMCBSCi9JRCBbPDQ0MTY0NzEwQjAzN0U5NkQxMzUwMzU2NjIwNTYyMjNBPjw0NDE2NDcxMEIwMzdFOTZEMTM1MDM1NjYyMDU2MjIzQT5dCj4+CnN0YXJ0eHJlZgo2NzEKJSVFT0YK"
)


def _hello_pdf_bytes() -> bytes:
    return base64.b64decode(_HELLO_PDF_B64)


def _docx_bytes(text: str) -> bytes:
    buf = BytesIO()
    doc = Document()
    doc.add_paragraph(text)
    doc.save(buf)
    return buf.getvalue()


def test_extract_txt() -> None:
    r = extract_document_bytes(
        filename="n.txt",
        mime="text/plain",
        raw=b"hello world",
    )
    assert r.status == "extracted"
    assert r.text == "hello world"


def test_extract_md() -> None:
    r = extract_document_bytes(
        filename="x.md",
        mime="text/markdown",
        raw=b"# Hi\n",
    )
    assert r.status == "extracted"
    assert "# Hi" in r.text


def test_extract_txt_truncates() -> None:
    big = "x" * (MAX_EXTRACTED_CHARS_PER_FILE + 500)
    r = extract_document_bytes(filename="big.txt", mime="text/plain", raw=big.encode())
    assert r.status == "truncated"
    assert len(r.text) == MAX_EXTRACTED_CHARS_PER_FILE


def test_extract_pdf_text() -> None:
    r = extract_document_bytes(
        filename="h.pdf",
        mime="application/pdf",
        raw=_hello_pdf_bytes(),
    )
    assert r.status == "extracted"
    assert "Hello PDF Line" in r.text


def test_extract_pdf_invalid_is_failed() -> None:
    r = extract_document_bytes(
        filename="bad.pdf",
        mime="application/pdf",
        raw=b"not a pdf",
    )
    assert r.status == "failed"
    assert r.error_reason == "invalid_pdf"


def test_extract_docx() -> None:
    raw = _docx_bytes("DOCX secret line")
    r = extract_document_bytes(
        filename="w.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        raw=raw,
    )
    assert r.status == "extracted"
    assert "DOCX secret line" in r.text


def test_extract_doc_unsupported() -> None:
    r = extract_document_bytes(
        filename="legacy.doc",
        mime="application/msword",
        raw=b"\xd0\xcf\x11\xe0",  # ole header-ish
    )
    assert r.status == "unsupported"


def test_total_message_budget_truncates_second_file() -> None:
    """First body uses budget; second is truncated to remainder; third is omitted."""
    r1 = extract_document_bytes(filename="1.txt", mime="text/plain", raw=b"X" * 1_000)
    r2 = extract_document_bytes(filename="2.txt", mime="text/plain", raw=b"Y" * 5_000)
    r3 = extract_document_bytes(filename="3.txt", mime="text/plain", raw=b"Z" * 100)
    sections = build_document_llm_sections([r1, r2, r3], max_total_chars=2_500)
    joined = "\n".join(sections)
    assert joined.count("X") == 1_000
    assert joined.count("Y") == 1_500
    assert joined.count("Z") == 0
    assert "3.txt" in joined


def test_build_sections_placeholder_no_budget_use() -> None:
    r = DocumentExtractionResult(
        filename="x.doc",
        mime="application/msword",
        status="unsupported",
        text="",
        truncated=False,
        error_reason=None,
    )
    sections = build_document_llm_sections([r], max_total_chars=0)
    assert len(sections) == 1
    assert "not enabled" in sections[0]


def test_format_block_no_gs_path() -> None:
    r = extract_document_bytes(filename="n.txt", mime="text/plain", raw=b"body")
    block = format_document_block_for_llm(r, content_body=r.text)
    assert "gs://" not in block
    assert "C:\\" not in block


def _small_xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    buf = BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "January"
    ws.append(["Agent", "Business", "Amount"])
    ws.append(["Ada", "North", "42"])
    wb.save(buf)
    return buf.getvalue()


def test_extract_xlsx_values_only() -> None:
    raw = _small_xlsx_bytes()
    r = extract_document_bytes(
        filename="commissions.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        raw=raw,
    )
    assert r.status == "extracted"
    assert "January" in r.text
    assert "Agent" in r.text and "Amount" in r.text
    assert "Ada" in r.text
    blk = format_document_block_for_llm(r, content_body=r.text)
    assert "[Attached spreadsheet:" in blk
    assert "gs://" not in blk


def test_extract_csv() -> None:
    r = extract_document_bytes(
        filename="rows.csv",
        mime="text/csv",
        raw=b"a,b\n1,2\n3,4\n",
    )
    assert r.status == "extracted"
    assert "Columns:" in r.text and "a" in r.text


def test_extract_csv_truncates_rows() -> None:
    lines = ["c,d"] + [f"{i},{i}" for i in range(150)]
    body = "\n".join(lines).encode()
    r = extract_document_bytes(filename="big.csv", mime="text/csv", raw=body)
    assert r.status == "truncated"


def test_extract_legacy_xls_unsupported() -> None:
    r = extract_document_bytes(
        filename="legacy.xls",
        mime="application/vnd.ms-excel",
        raw=b"\xd0\xcf\x11\xe0",
    )
    assert r.status == "unsupported"
    sec = build_document_llm_sections([r], max_total_chars=1000)[0]
    assert "legacy .xls extraction is not enabled" in sec
    assert "gs://" not in sec

